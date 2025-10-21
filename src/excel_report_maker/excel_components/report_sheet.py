from dataclasses import dataclass, field

import pandas as pd

from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from excel_report_maker.excel_components.report_table import ReportTable
from typing import Optional, TYPE_CHECKING
if TYPE_CHECKING:
    from excel_report_maker.excel_reporter.excel_report import ExcelReport

@dataclass
class ReportSheetOptions:
    column_padding: int = 4
    gridlines: bool = True


@dataclass
class ReportSheet:
    sheet_name: Optional[str] = None
    tables: list["ReportTable"] = field(default_factory=list)
    settings: ReportSheetOptions = field(default_factory=ReportSheetOptions)

    @property
    def ws(self) -> Optional[Worksheet]:
        return getattr(self, '_ws', None)

    @staticmethod
    def from_table_list(sheet_name: str, tables: list["ReportTable"], *args, **kwargs):
        return ReportSheet(sheet_name=sheet_name, tables=tables, *args, **kwargs)
    
    def register_table(self, table: "ReportTable"):
        if not isinstance(table, ReportTable):
            raise TypeError("`table` must be a ReportTable.")
        self.tables.append(table)
    
    def _create_report_sheet(self, excel_report: "ExcelReport"):
        # Create and save a reference to the worksheet
        ws = excel_report.wb.create_sheet(self.sheet_name)
        self._ws = ws
        current_row = 1

        # Write each query result (table) from the file into the sheet.
        for table in self.tables:
            current_row = table._append_table_to_sheet(
                report_sheet=self,
                global_table_counter=excel_report.global_table_counter,
                start_row=current_row
            )
            excel_report.global_table_counter += 1

        # Auto-adjust column widths.
        for col in ws.columns:
            max_length = max((len(str(cell.value)) if cell.value else 0 for cell in col), default=0)
            adjusted_width = max_length + self.settings.column_padding
            ws.column_dimensions[get_column_letter(col[0].column)].width = adjusted_width
        
        # Optionally disable gridlines
        ws.sheet_view.showGridLines = self.settings.gridlines